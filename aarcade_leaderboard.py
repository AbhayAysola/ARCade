#!/usr/bin/env python3
"""Arcade leaderboards backed by an editable Excel workbook.

Keep leaderboard.xlsx beside this script. Edit the Players sheet and save it;
the web page reloads scores automatically.
"""
import json
import os
import sys
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

try:
    from openpyxl import load_workbook, Workbook
except ImportError:
    load_workbook = Workbook = None

PORT = 5000
TOP_N = 15
TITLE = "ARCADE LEADERBOARDS"
WORKBOOK = os.path.join(os.path.dirname(os.path.abspath(__file__)), "leaderboard.xlsx")
GAMES = ["Tetris", "Pac Man", "Birds with Guns"]


def read_scores():
    """Read current player scores from the Excel workbook on each API request."""
    if load_workbook is None:
        return {"error": "Missing dependency: install openpyxl with: python -m pip install openpyxl",
                "games": {game: [] for game in GAMES}, "updated": time.strftime("%H:%M:%S")}
    if not os.path.exists(WORKBOOK):
        return {"error": "Excel file not found: leaderboard.xlsx", "games": {game: [] for game in GAMES},
                "updated": time.strftime("%H:%M:%S")}
    try:
        wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
        if "Players" not in wb.sheetnames:
            wb.close()
            return {"error": 'The workbook needs a sheet named "Players".',
                    "games": {game: [] for game in GAMES}, "updated": time.strftime("%H:%M:%S")}
        ws = wb["Players"]
        values = ws.iter_rows(values_only=True)
        headers = next(values, ())
        col_map = {str(v).strip().lower(): i for i, v in enumerate(headers) if v is not None}
        player_col = col_map.get("player")
        game_cols = {game: col_map.get(game.lower()) for game in GAMES}
        if player_col is None or any(i is None for i in game_cols.values()):
            wb.close()
            return {"error": "Headers must be: Player, Tetris, Pac Man, Birds with Guns.",
                    "games": {game: [] for game in GAMES}, "updated": time.strftime("%H:%M:%S")}
        results = {game: [] for game in GAMES}
        for row in values:
            if player_col >= len(row) or row[player_col] is None or not str(row[player_col]).strip():
                continue
            player = str(row[player_col]).strip()
            for game, col in game_cols.items():
                raw = row[col] if col < len(row) else 0
                try:
                    score = float(raw or 0)
                    if score.is_integer():
                        score = int(score)
                except (TypeError, ValueError):
                    score = 0
                results[game].append({"player": player, "score": score})
        wb.close()
        for game in GAMES:
            results[game].sort(key=lambda r: (-r["score"], r["player"].casefold()))
            results[game] = results[game][:TOP_N]
        return {"error": "", "games": results, "updated": time.strftime("%H:%M:%S")}
    except Exception as exc:
        return {"error": "Could not read Excel file: " + str(exc),
                "games": {game: [] for game in GAMES}, "updated": time.strftime("%H:%M:%S")}


PAGE = r"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Arcade Leaderboards</title>
<style>
*{box-sizing:border-box}html,body{min-height:100%;margin:0}
body{background:#0b0620;color:#fff;text-align:center;font-family:'Courier New',monospace;font-weight:bold;padding:20px 12px 40px}
body:before{content:"";position:fixed;inset:0;z-index:-1;background:linear-gradient(#0b0620,#1b0a45 60%,#2f0f66)}
h1{font-size:clamp(1.6rem,5vw,3rem);color:#fff078;text-shadow:3px 3px #ff2fa8,0 0 20px #ff3cbe;margin:12px 0}
#sub{color:#6af0ff;margin:12px}
.tabs{display:flex;gap:8px;justify-content:center;flex-wrap:wrap;margin:24px auto 14px}
button{font: bold 1rem 'Courier New',monospace;padding:12px 15px;border:2px solid #a78bfa;border-radius:7px;background:#1c1042;color:#fff;cursor:pointer}
button.active{background:#ff3cbe;color:#160626;border-color:#ffd34d}
#status{min-height:1.5em;color:#c8b8ff;font-size:.85rem;margin:10px}
table{width:min(900px,96%);margin:18px auto;border-collapse:separate;border-spacing:0 6px;table-layout:fixed}
td,th{padding:12px 10px;background:rgba(20,8,60,.88);overflow-wrap:anywhere}
th{color:#6af0ff}td:first-child,th:first-child{width:18%}td:nth-child(2),th:nth-child(2){text-align:left}
td:last-child,th:last-child{width:30%;text-align:right;color:#fff078}
tr:first-child td{background:rgba(255,60,190,.22)}tr:first-child td:first-child{color:#ffd34d}
#empty{color:#c8b8ff;margin:45px 0}
#updated{color:#9b8fc9;font-size:.75rem}
</style></head><body>
<h1 id="title">ARCADE LEADERBOARDS</h1>
<div id="sub">THREE GAMES. THREE CHANCES TO BE NUMBER ONE.</div>
<div class="tabs" id="tabs"></div>
<div id="status">Loading scores from Excel...</div>
<div id="list"></div>
<div id="updated"></div>
<script>
const games = ["Tetris", "Pac Man", "Birds with Guns"];
let selected = games[0], latest = null;
const tabs = document.getElementById('tabs');
games.forEach(game => {
  const b = document.createElement('button');
  b.textContent = game;
  b.onclick = () => { selected = game; render(latest); };
  b.dataset.game = game;
  tabs.appendChild(b);
});
function el(tag, cls, text) {
  const e = document.createElement(tag);
  if (cls) e.className = cls;
  if (text !== undefined) e.textContent = text;
  return e;
}
function render(data) {
  if (!data) return;
  latest = data;
  document.querySelectorAll('button').forEach(b => b.classList.toggle('active', b.dataset.game === selected));
  const list = document.getElementById('list');
  list.innerHTML = '';
  document.getElementById('status').textContent = data.error || ('LEADERBOARD: ' + selected.toUpperCase());
  document.getElementById('updated').textContent = 'Last refreshed: ' + data.updated + ' | Edit leaderboard.xlsx to update scores';
  const rows = (data.games && data.games[selected]) || [];
  if (!rows.length) {
    list.appendChild(el('div', 'empty', data.error ? 'Check the Excel file and restart if needed.' : 'No players entered yet. Add players and scores to the Excel sheet.'));
    return;
  }
  const table = document.createElement('table');
  const head = document.createElement('tr');
  ['RANK', 'PLAYER', 'SCORE'].forEach(t => head.appendChild(el('th', '', t)));
  table.appendChild(head);
  rows.forEach((row, i) => {
    const tr = document.createElement('tr');
    tr.appendChild(el('td', '', i === 0 ? '1st' : i === 1 ? '2nd' : i === 2 ? '3rd' : String(i + 1)));
    tr.appendChild(el('td', '', row.player));
    tr.appendChild(el('td', '', Number(row.score).toLocaleString()));
    table.appendChild(tr);
  });
  list.appendChild(table);
}
async function load() {
  try {
    const response = await fetch('/api/top?t=' + Date.now(), {cache:'no-store'});
    render(await response.json());
  } catch (e) {
    document.getElementById('status').textContent = 'Cannot load scores. Is the leaderboard server running?';
  }
}
load(); setInterval(load, 5000);
</script></body></html>"""


class Handler(BaseHTTPRequestHandler):
    def send_body(self, code, body, content_type):
        body = body.encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        path = self.path.split("?")[0]
        if path == "/api/top":
            self.send_body(200, json.dumps(read_scores()), "application/json; charset=utf-8")
        else:
            self.send_body(200, PAGE, "text/html; charset=utf-8")

    def log_message(self, *args):
        pass


if __name__ == "__main__":
    if load_workbook is None:
        print("Install dependency first: python -m pip install openpyxl")
        sys.exit(1)
    if not os.path.exists(WORKBOOK):
        wb = Workbook()
        ws = wb.active
        ws.title = "Players"
        ws.append(["Player", "Tetris", "Pac Man", "Birds with Guns"])
        ws.append(["Example Player", 1200, 850, 2400])
        wb.save(WORKBOOK)
        print("Created starter Excel workbook:", WORKBOOK)
    print("Arcade leaderboards at http://localhost:%d" % PORT)
    print("Edit %s to update player scores." % os.path.basename(WORKBOOK))
    ThreadingHTTPServer(("0.0.0.0", PORT), Handler).serve_forever()

